"""Text extraction from uploaded documents, run in an isolated, time-limited child process.

Untrusted documents are parsed in a separate process so a pathological file can't hang or crash
the API. The child is killed on timeout. Output is plain text plus page start offsets, so
evidence spans (char offsets) can be mapped back to page numbers.
"""

from __future__ import annotations

import io
import multiprocessing as mp
import re
from dataclasses import dataclass, field
from multiprocessing.connection import Connection
from typing import Any

from nexscout.domain.errors import ValidationFailed
from nexscout.infrastructure.ingestion.files import FileKind

MAX_PAGES = 60
MAX_TEXT_CHARS = 400_000
DEFAULT_TIMEOUT_S = 20.0


@dataclass(frozen=True)
class ParsedText:
    text: str
    page_starts: list[int] = field(default_factory=list)  # char offset where each page begins

    @property
    def page_count(self) -> int | None:
        return len(self.page_starts) or None

    def page_of(self, offset: int) -> int | None:
        """1-based page number containing `offset`, or None for unpaginated formats."""
        if not self.page_starts:
            return None
        page = 0
        for i, start in enumerate(self.page_starts):
            if start <= offset:
                page = i
            else:
                break
        return page + 1


def normalize_text(raw: str) -> str:
    text = raw.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    text = re.sub(r"[ \t ]+\n", "\n", text)  # trailing spaces
    text = re.sub(r"\n{4,}", "\n\n\n", text)
    return text


def _parse_pdf(data: bytes) -> ParsedText:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise ValidationFailed("This PDF is password-protected. Upload an unprotected copy.")
        if len(reader.pages) > MAX_PAGES:
            raise ValidationFailed(f"This PDF has more than {MAX_PAGES} pages.")
        parts: list[str] = []
        starts: list[int] = []
        offset = 0
        for page in reader.pages:
            page_text = normalize_text(page.extract_text() or "").strip("\n")
            starts.append(offset)
            parts.append(page_text)
            offset += len(page_text) + 2  # "\n\n" separator
        text = "\n\n".join(parts)
    except PdfReadError as exc:
        raise ValidationFailed("This PDF could not be read. It may be damaged.") from exc
    if not text.strip():
        raise ValidationFailed(
            "No text found in this PDF. It may be a scanned image; OCR isn't supported yet."
        )
    return ParsedText(text=text, page_starts=starts)


def _parse_docx(data: bytes) -> ParsedText:
    import docx

    document = docx.Document(io.BytesIO(data))
    lines: list[str] = []
    for paragraph in document.paragraphs:
        lines.append(paragraph.text)
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                lines.append(" | ".join(cells))
    return ParsedText(text=normalize_text("\n".join(lines)).strip("\n"))


def _parse_text(data: bytes) -> ParsedText:
    return ParsedText(text=normalize_text(data.decode("utf-8")).strip("\n"))


def parse_bytes(data: bytes, kind: FileKind) -> ParsedText:
    """Parse in-process. Prefer `parse_isolated` for untrusted uploads."""
    if kind is FileKind.PDF:
        parsed = _parse_pdf(data)
    elif kind is FileKind.DOCX:
        parsed = _parse_docx(data)
    else:
        parsed = _parse_text(data)
    if len(parsed.text) > MAX_TEXT_CHARS:
        raise ValidationFailed("This document contains too much text to process.")
    if not parsed.text.strip():
        raise ValidationFailed("No readable text found in this document.")
    return parsed


def _child(conn: Connection, data: bytes, kind: str) -> None:
    try:
        parsed = parse_bytes(data, FileKind(kind))
        conn.send(("ok", parsed.text, parsed.page_starts))
    except ValidationFailed as exc:
        conn.send(("invalid", exc.message, None))
    except Exception as exc:  # noqa: BLE001 - report any parser crash as a typed failure
        conn.send(("error", type(exc).__name__, None))
    finally:
        conn.close()


def parse_isolated(data: bytes, kind: FileKind, timeout_s: float = DEFAULT_TIMEOUT_S) -> ParsedText:
    """Parse in a spawned child process; kill it if it exceeds `timeout_s`. Blocking call."""
    ctx = mp.get_context("spawn")
    parent, child = ctx.Pipe(duplex=False)
    proc = ctx.Process(target=_child, args=(child, data, kind.value), daemon=True)
    proc.start()
    child.close()
    try:
        if not parent.poll(timeout_s):
            raise ValidationFailed("This document took too long to read and was skipped.")
        result: tuple[str, Any, Any] = parent.recv()
    except EOFError as exc:
        raise ValidationFailed("This document could not be read.") from exc
    finally:
        parent.close()
        if proc.is_alive():
            proc.kill()
        proc.join(timeout=5)
    status, payload, starts = result
    if status == "ok":
        return ParsedText(text=payload, page_starts=list(starts))
    if status == "invalid":
        raise ValidationFailed(payload)
    raise ValidationFailed("This document could not be read.")
