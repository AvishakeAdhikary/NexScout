"""Render canonical Markdown to HTML, PDF (headless Chromium) and DOCX (python-docx), then
validate the files by reading them back.

- Raw HTML in Markdown is escaped (markdown-it `html=False`); unsafe link schemes are refused.
- PDF rendering runs Chromium with JavaScript off and every network request aborted.
- Placeholders stay visible and highlighted in every format, so an unfinished document can't be
  mistaken for a finished one.
- Playwright's sync API runs in a worker thread: the worker's Selector event loop (needed by
  psycopg on Windows) can't start subprocesses, while a fresh thread gets its own loop.
"""

from __future__ import annotations

import asyncio
import html as html_lib
import io
import re
import zipfile
from typing import Any

from docx import Document as DocxDocument
from docx.enum.text import WD_COLOR_INDEX
from docx.opc.exceptions import PackageNotFoundError
from docx.text.paragraph import Paragraph
from markdown_it import MarkdownIt
from markdown_it.token import Token
from pypdf import PdfReader
from pypdf.errors import PyPdfError

from nexscout.domain.drafting import PLACEHOLDER, find_placeholders

_MD = MarkdownIt("commonmark", {"html": False, "linkify": False, "typographer": False})

_CSS = """
@page { size: __PAPER__; margin: 22mm 20mm; }
body { font-family: "Georgia", "Times New Roman", serif; font-size: 11pt; line-height: 1.45;
       color: #111; }
h1 { font-size: 20pt; margin: 0 0 6pt; }
h2 { font-size: 13pt; margin: 14pt 0 4pt; border-bottom: 1px solid #999; padding-bottom: 2pt; }
h3 { font-size: 11.5pt; margin: 10pt 0 2pt; }
p { margin: 0 0 7pt; }
ul, ol { margin: 0 0 7pt 16pt; padding: 0; }
li { margin: 0 0 3pt; }
a { color: #111; text-decoration: none; }
mark.ph { background: #ffe066; padding: 0 2px; }
"""


def markdown_to_html(markdown: str, title: str, paper: str = "A4") -> str:
    body = _MD.render(markdown)
    css = _CSS.replace("__PAPER__", "letter" if paper == "Letter" else "A4")
    body = PLACEHOLDER.sub(lambda m: f'<mark class="ph">{m.group(0)}</mark>', body)
    return (
        "<!doctype html><html><head><meta charset='utf-8'>"
        f"<title>{html_lib.escape(title)}</title>"
        f"<style>{css}</style></head>"
        f"<body>{body}</body></html>"
    )


def _pdf_sync(html: str) -> bytes:
    from playwright.sync_api import sync_playwright

    with sync_playwright() as p:
        browser = p.chromium.launch(headless=True)
        try:
            context = browser.new_context(java_script_enabled=False, offline=True)
            context.route("**/*", lambda route: route.abort())
            page = context.new_page()
            page.set_content(html, wait_until="load")
            return page.pdf(format="A4", print_background=True, prefer_css_page_size=True)
        finally:
            browser.close()


async def html_to_pdf(html: str) -> bytes:
    return await asyncio.to_thread(_pdf_sync, html)


def _add_runs(par: Paragraph, children: list[Token]) -> None:
    bold = italic = False
    link: str | None = None
    for t in children:
        match t.type:
            case "strong_open":
                bold = True
            case "strong_close":
                bold = False
            case "em_open":
                italic = True
            case "em_close":
                italic = False
            case "link_open":
                link = str(t.attrs.get("href") or "")
            case "link_close":
                if link and link.startswith(("http://", "https://", "mailto:")):
                    par.add_run(f" ({link})")
                link = None
            case "softbreak":
                par.add_run(" ")
            case "hardbreak":
                par.add_run().add_break()
            case "text" | "code_inline":
                pos = 0
                for m in PLACEHOLDER.finditer(t.content):
                    if m.start() > pos:
                        run = par.add_run(t.content[pos : m.start()])
                        run.bold, run.italic = bold or None, italic or None
                    ph = par.add_run(m.group(0))
                    ph.font.highlight_color = WD_COLOR_INDEX.YELLOW
                    pos = m.end()
                if pos < len(t.content):
                    run = par.add_run(t.content[pos:])
                    run.bold, run.italic = bold or None, italic or None


def markdown_to_docx(markdown: str, title: str, paper: str = "A4") -> bytes:
    from docx.shared import Cm, Inches

    doc = DocxDocument()
    for section in doc.sections:  # page size follows the regional format
        if paper == "Letter":
            section.page_width, section.page_height = Inches(8.5), Inches(11)
        else:
            section.page_width, section.page_height = Cm(21.0), Cm(29.7)
    doc.core_properties.title = title[:250]
    doc.core_properties.author = ""
    doc.core_properties.comments = "Created with NexScout"
    tokens = _MD.parse(markdown)
    lists: list[str] = []
    i = 0
    while i < len(tokens):
        t = tokens[i]
        if t.type == "heading_open":
            level = int(t.tag[1])
            inline = tokens[i + 1]
            text = inline.content
            if level == 1:
                doc.add_heading(text, level=0)
            else:
                doc.add_heading(text, level=min(level - 1, 3))
            i += 3
            continue
        if t.type in ("bullet_list_open", "ordered_list_open"):
            lists.append("List Bullet" if t.type == "bullet_list_open" else "List Number")
        elif t.type in ("bullet_list_close", "ordered_list_close"):
            lists.pop()
        elif t.type == "paragraph_open":
            inline = tokens[i + 1]
            style = lists[-1] if lists else None
            par = doc.add_paragraph(style=style) if style else doc.add_paragraph()
            _add_runs(par, inline.children or [])
            i += 3
            continue
        elif t.type == "hr":
            doc.add_paragraph("")
        i += 1
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _norm(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip().lower()


def headings(markdown: str) -> list[str]:
    return [m.group(1).strip() for m in re.finditer(r"^#{1,3}\s+(.+)$", markdown, re.M)]


def validate_pdf(data: bytes, markdown: str) -> dict[str, Any]:
    checks: dict[str, Any] = {"opens": False}
    try:
        reader = PdfReader(io.BytesIO(data))
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
    except (PyPdfError, ValueError, OSError, KeyError) as exc:
        checks["error"] = f"Couldn't read the PDF back: {type(exc).__name__}"
        return checks
    flat = _norm(text)
    expected = headings(markdown)
    missing = [h for h in expected if _norm(h) not in flat]
    checks.update(
        opens=True,
        pages=len(reader.pages),
        headings_found=len(expected) - len(missing),
        headings_missing=missing,
        placeholders=len(find_placeholders(text.replace("\n", " "))),
    )
    checks["ok"] = checks["pages"] > 0 and not missing
    return checks


def validate_docx(data: bytes, markdown: str) -> dict[str, Any]:
    checks: dict[str, Any] = {"opens": False}
    try:
        doc = DocxDocument(io.BytesIO(data))
    except (PackageNotFoundError, zipfile.BadZipFile, KeyError, ValueError, OSError) as exc:
        checks["error"] = f"Couldn't read the Word file back: {type(exc).__name__}"
        return checks
    text = "\n".join(p.text for p in doc.paragraphs)
    flat = _norm(text)
    expected = headings(markdown)
    missing = [h for h in expected if _norm(h) not in flat]
    checks.update(
        opens=True,
        paragraphs=len(doc.paragraphs),
        headings_found=len(expected) - len(missing),
        headings_missing=missing,
        placeholders=len(find_placeholders(text)),
    )
    checks["ok"] = checks["paragraphs"] > 0 and not missing
    return checks
