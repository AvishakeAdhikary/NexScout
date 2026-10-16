"""Rendering: escaping, placeholders kept visible, PDF/DOCX produced and validated."""

import io

from docx import Document

from nexscout.infrastructure.render import (
    html_to_pdf,
    markdown_to_docx,
    markdown_to_html,
    validate_docx,
    validate_pdf,
)

MD = """# Statement of purpose

## Motivation

I want to study NLP. [[PLACEHOLDER: why this group]]

## Background

- **MSc**, Sampleford
- Python

<script>alert(1)</script> [click](javascript:alert(1)) [site](https://uni.example/x)
"""


def test_html_is_escaped_and_placeholders_marked() -> None:
    html = markdown_to_html(MD, "SOP <b>")
    assert "<script>" not in html
    assert "&lt;script&gt;" in html
    assert 'href="javascript' not in html
    assert '<mark class="ph">[[PLACEHOLDER: why this group]]</mark>' in html
    assert "<title>SOP &lt;b&gt;</title>" in html


def test_docx_structure_and_highlight() -> None:
    data = markdown_to_docx(MD, "SOP")
    doc = Document(io.BytesIO(data))
    styles = [(p.style.name if p.style else "", p.text) for p in doc.paragraphs]
    assert ("Title", "Statement of purpose") in styles
    assert ("Heading 1", "Motivation") in styles
    assert any(s == "List Bullet" and t == "MSc, Sampleford" for s, t in styles)
    highlighted = [r.text for p in doc.paragraphs for r in p.runs if r.font.highlight_color]
    assert highlighted == ["[[PLACEHOLDER: why this group]]"]
    assert any("https://uni.example/x" in t for _, t in styles)
    assert validate_docx(data, MD)["ok"] is True


async def test_pdf_renders_offline_and_validates() -> None:
    pdf = await html_to_pdf(markdown_to_html(MD, "SOP"))
    assert pdf.startswith(b"%PDF")
    checks = validate_pdf(pdf, MD)
    assert checks["ok"] is True
    assert checks["pages"] == 1
    assert checks["placeholders"] == 1


def test_validators_catch_broken_or_incomplete_files() -> None:
    assert validate_pdf(b"not a pdf", MD)["opens"] is False
    assert validate_docx(b"not a docx", MD)["opens"] is False
    other = markdown_to_docx("# Something else\n", "x")
    checks = validate_docx(other, MD)
    assert checks["ok"] is False
    assert "Motivation" in checks["headings_missing"]
