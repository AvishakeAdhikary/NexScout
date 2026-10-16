"""Builders for synthetic test documents (no binary fixtures committed)."""

from __future__ import annotations

import io
from pathlib import Path

FIXTURES = Path(__file__).parent
ADA_CV_MD = FIXTURES / "personas" / "ada_example_cv.md"


def ada_cv_text() -> str:
    return ADA_CV_MD.read_text(encoding="utf-8")


def make_docx(lines: list[str]) -> bytes:
    import docx

    document = docx.Document()
    for line in lines:
        if line.startswith("## "):
            document.add_heading(line[3:], level=2)
        elif line.startswith("# "):
            document.add_heading(line[2:], level=1)
        else:
            document.add_paragraph(line)
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def _pdf_escape(text: str) -> str:
    return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def make_pdf(pages: list[list[str]]) -> bytes:
    """Minimal valid PDF (Helvetica, ASCII text) with one page per list of lines."""
    objects: list[bytes] = []
    page_ids: list[int] = []
    # 1: catalog, 2: pages, 3: font; pages and contents follow.
    font_id = 3
    next_id = 4
    page_objs: list[tuple[int, int, bytes]] = []
    for lines in pages:
        content_lines = ["BT", "/F1 11 Tf", "50 780 Td", "14 TL"]
        for line in lines:
            ascii_line = line.encode("ascii", "replace").decode("ascii")
            content_lines.append(f"({_pdf_escape(ascii_line)}) Tj T*")
        content_lines.append("ET")
        stream = "\n".join(content_lines).encode("ascii")
        page_id, content_id = next_id, next_id + 1
        next_id += 2
        page_ids.append(page_id)
        page_objs.append((page_id, content_id, stream))

    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    kids = " ".join(f"{pid} 0 R" for pid in page_ids)
    objects.append(f"<< /Type /Pages /Kids [{kids}] /Count {len(page_ids)} >>".encode())
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    body: dict[int, bytes] = {1: objects[0], 2: objects[1], font_id: objects[2]}
    for page_id, content_id, stream in page_objs:
        body[page_id] = (
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 842] "
            f"/Resources << /Font << /F1 {font_id} 0 R >> >> /Contents {content_id} 0 R >>"
        ).encode()
        body[content_id] = (
            f"<< /Length {len(stream)} >>\nstream\n".encode() + stream + b"\nendstream"
        )

    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets: dict[int, int] = {}
    for obj_id in sorted(body):
        offsets[obj_id] = out.tell()
        out.write(f"{obj_id} 0 obj\n".encode() + body[obj_id] + b"\nendobj\n")
    xref = out.tell()
    count = max(body) + 1
    out.write(f"xref\n0 {count}\n0000000000 65535 f \n".encode())
    for obj_id in range(1, count):
        out.write(f"{offsets[obj_id]:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {count} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return out.getvalue()
